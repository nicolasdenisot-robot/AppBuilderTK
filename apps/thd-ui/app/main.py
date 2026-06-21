"""
thd-ui — THD Calculator with a Tkinter interface.

The user selects which NVGate window to use and how many harmonics to include,
then clicks Run.  Results appear in NVGate (Harmonic marker + scalar window).
After a successful run, "Export to Excel" creates a 3-sheet report:
  Sheet 1 "Metadata"    — full context: project, signal, FFT parameters, NVGate info
  Sheet 2 "Spectrum"    — frequency/amplitude table + scatter chart with harmonic markers
  Sheet 3 "THD Analysis"— harmonic table, formula, result
"""

import sys
import math
import datetime
import threading
from pathlib import Path
import tkinter as tk
from tkinter import ttk, filedialog

# ── Path setup: works both inside this workspace and in an exported bundle ──
_HERE = Path(__file__).resolve()
for _candidate in (_HERE.parents[3] / "shared", _HERE.parents[1]):
    if (_candidate / "oros_helpers").is_dir() and str(_candidate) not in sys.path:
        sys.path.insert(0, str(_candidate))

import oros_helpers  # noqa: F401  — side-effect: calls ensure_pynvdrive()
from oros_helpers import connect

import pynvdrive
from pynvdrive.command import CommandAction
from pynvdrive.commands.marker import GetWindowTraces, GetWindowDisplayZone
from pynvdrive.commands.graphics import AddWindowResult
from pynvdrive.commands.graphics.getcurrentlayout import GetCurrentLayout
from pynvdrive.commands.graphics.getwindownamelist import GetWindowNameList
from pynvdrive.commands.signalresult import (
    CreateResultEx, GetResultEx, DeleteAllResultEx, CreateTCPResultChannel,
)
from pynvdrive.commands.project.getcurrentprojectname import GetCurrentProjectName
from pynvdrive.commands.project.getdatabasepath import GetDatabasePath
from pynvdrive.commands.analyzers.getnvgateinfo import GetNVGateInfo
from pynvdrive.formats.result import Result
from pynvdrive.formats.blocks import HEAD, SCAL, INFO, TCPRESULT

# ── Lookup tables for coded values in the result blocks ────────────────────
_RESULT_TYPE = {1: "Scalar", 2: "Spectrum (regular vector)",
                3: "Reference vector", 8: "Waterfall"}
_MEASURE_STATUS = {0: "No measurement", 1: "In progress", 2: "Finished"}
_N_OCTAVE = {0: "Linear (FFT)", 1: "1/1 octave", 2: "1/3 octave",
             3: "1/12 octave", 4: "1/24 octave"}
_WEIGHTING_WINDOW = {
    0: "Rectangular (none)", 1: "Hann", 2: "Flat-top",
    3: "Hamming", 4: "Blackman", 5: "Exponential",
    6: "Force", 7: "Exponential + Force", 8: "Kaiser-Bessel",
}


# ── Custom NVDrive command (not yet in pynvdrive Python library) ────────────

class AddHarmonicMarker(CommandAction):
    """Add a harmonic marker on a display zone (NVDrive manual §3.2.9)."""
    def __init__(self, window_name, start_position, num_harmonics=5, display_zone=''):
        super().__init__(
            name=None,
            parameters=[window_name, str(start_position), str(num_harmonics), str(display_zone)],
        )


# ── THD mathematics ─────────────────────────────────────────────────────────

def get_amplitudes(result):
    """Return a flat list of linear amplitude values from a spectrum result."""
    if result.rvec is not None:
        return list(result.rvec.data)
    if result.cvec is not None:
        return [math.sqrt(d[0] ** 2 + d[1] ** 2) for d in result.cvec.data]
    return None


def find_fundamental(amplitudes, freq_start, freq_stop):
    """Return (index, frequency_Hz, amplitude) of the highest-amplitude bin."""
    idx = max(range(len(amplitudes)), key=lambda i: amplitudes[i])
    n = len(amplitudes)
    freq = freq_start + idx * (freq_stop - freq_start) / max(n - 1, 1)
    return idx, freq, amplitudes[idx]


def amplitude_at(amplitudes, freq_start, freq_stop, target_hz):
    """Amplitude at the bin nearest to target_hz."""
    n = len(amplitudes)
    df = (freq_stop - freq_start) / max(n - 1, 1)
    idx = int(round((target_hz - freq_start) / df))
    idx = max(0, min(idx, n - 1))
    return amplitudes[idx]


def compute_thd(h1, overtones):
    """THD (%) = sqrt(H2^2+H3^2+...+Hn^2) / |H1| x 100.  overtones = [H2, H3, ...]."""
    if abs(h1) == 0:
        raise ValueError("Fundamental amplitude is zero — cannot compute THD.")
    return math.sqrt(sum(h ** 2 for h in overtones)) / abs(h1) * 100.0


def inject_scalar_percent(client, value_pct, name="THD"):
    """Push a percentage scalar into NVGate and return the TCP channel id."""
    head = HEAD(
        version=5, result_type=1, size=1, dimension=1, is_complex=0,
        measure_status=2,
        physical_quantity="", unit_name=None, unit_label=None,
        name=name, overall_level_status=0, weighted_ol_status=0,
        weighting_window=0, magnitude_key="Percentage",
    )
    scal = SCAL(version=2, min=value_pct, max=value_pct, value=value_pct,
                rel_precision=0.001, abs_precision=1e-30, imag=0.0)
    info = INFO(version=1, info_mask=0, date=datetime.datetime.now(),
                global_level=0, weighted_global_level=0, tacho=0)
    result = Result(head=head, scal=scal, info=info)
    tcp_res = TCPRESULT(module=498, process=182, channel=1, ref_channel=0,
                        input_name="THD_Result", result=result)
    cmd = CreateTCPResultChannel(tcp_result=tcp_res)
    client.send_command(cmd)
    return cmd.value


def _safe_get(obj, attr, default=''):
    """Return obj.attr if it exists and is not None; otherwise default."""
    val = getattr(obj, attr, None)
    return val if val is not None else default


def _decode(code, lookup, default_prefix='code'):
    """Map an integer code through a lookup dict, falling back to 'code N'."""
    if code is None:
        return ''
    return lookup.get(code, f'{default_prefix} {code}')


def _gather_metadata(client, result, trace, layout, window_name):
    """
    Collect every piece of context we can from the live NVGate session.
    All sub-commands are wrapped in try/except so a failure in one
    does not abort the export.
    Returns a flat dict of label → value pairs (all values are strings or numbers).
    """
    meta = {}

    # ── NVGate application info ─────────────────────────────────────────────
    try:
        cmd = GetNVGateInfo()
        client.send_command(cmd)
        meta['NVGate version']  = f"{cmd.nvgate_major_version}.{cmd.nvgate_minor_version}"
        meta['Build number']    = _safe_get(cmd, 'full_build_number')
        meta['Hardware type']   = _safe_get(cmd, 'hard_type')
        meta['Serial number']   = _safe_get(cmd, 'serial_number')
    except Exception:
        meta['NVGate version'] = meta['Build number'] = ''
        meta['Hardware type']  = meta['Serial number'] = ''

    # ── Project / data context ──────────────────────────────────────────────
    try:
        cmd = GetCurrentProjectName()
        client.send_command(cmd)
        meta['Project'] = cmd.value or '(not saved / live session)'
    except Exception:
        meta['Project'] = ''

    try:
        cmd = GetDatabasePath()
        client.send_command(cmd)
        meta['Database path'] = cmd.value or ''
    except Exception:
        meta['Database path'] = ''

    meta['Layout']        = layout
    meta['NVGate window'] = window_name

    # ── Signal / analysis identity ──────────────────────────────────────────
    meta['Signal (channel)'] = trace.input_name or ''
    try:
        meta['Analysis'] = trace.analysis_fullname
    except Exception:
        meta['Analysis'] = f"module {trace.module_id}, process {trace.process_id}"

    meta['Result name']      = _safe_get(result.head, 'name')
    meta['Physical quantity'] = _safe_get(result.head, 'physical_quantity')
    meta['Unit']             = (_safe_get(result.head, 'unit_label')
                                or _safe_get(result.head, 'unit_name'))
    meta['Magnitude type']   = _safe_get(result.head, 'magnitude_key')

    # Coded values — show human-readable label + raw code in parentheses
    rt   = _safe_get(result.head, 'result_type')
    ms   = _safe_get(result.head, 'measure_status')
    ww   = _safe_get(result.head, 'weighting_window')
    is_c = _safe_get(result.head, 'is_complex')

    meta['Result type']         = (_decode(rt,   _RESULT_TYPE,       'type')
                                   + (f'  ({rt})' if rt != '' else ''))
    meta['Measurement status']  = (_decode(ms,   _MEASURE_STATUS,    'status')
                                   + (f'  ({ms})' if ms != '' else ''))
    meta['Spectral window']     = (_decode(ww,   _WEIGHTING_WINDOW,  'code')
                                   + (f'  ({ww})' if ww != '' else ''))
    meta['Complex data']        = 'Yes' if is_c else 'No'

    # ── Measurement date & levels from INFO block ───────────────────────────
    if result.info is not None:
        d = _safe_get(result.info, 'date')
        try:
            meta['Measurement date'] = d.strftime('%Y-%m-%d  %H:%M:%S') if d else ''
        except Exception:
            meta['Measurement date'] = str(d)

        gl  = _safe_get(result.info, 'global_level', 0.0)
        wgl = _safe_get(result.info, 'weighted_global_level', 0.0)
        tch = _safe_get(result.info, 'tacho', 0.0)
        unit_lbl = meta['Unit'] or ''

        meta['Global RMS level']          = f"{gl:.6g}  {unit_lbl}" if gl else ''
        meta['A-weighted global level']   = f"{wgl:.6g}  {unit_lbl}" if wgl else ''
        meta['Tacho speed at measurement']= f"{tch:.4g} RPM"         if tch else ''
    else:
        meta['Measurement date'] = meta['Global RMS level'] = ''
        meta['A-weighted global level'] = meta['Tacho speed at measurement'] = ''

    # ── X-axis (frequency axis) from XREG ───────────────────────────────────
    if result.xreg is not None:
        no = _safe_get(result.xreg, 'n_octave')
        meta['Spectrum type'] = _decode(no, _N_OCTAVE, 'type') if no != '' else ''
        meta['X-axis unit']   = (_safe_get(result.xreg, 'unit_label')
                                 or _safe_get(result.xreg, 'unit_name'))
    else:
        meta['Spectrum type'] = meta['X-axis unit'] = ''

    # ── Export provenance ────────────────────────────────────────────────────
    meta['Exported by']  = 'THD Calculator  (thd-ui)'
    meta['Export date']  = datetime.datetime.now().strftime('%Y-%m-%d  %H:%M:%S')

    return meta


# ── NVGate helpers ──────────────────────────────────────────────────────────

DEFAULT_WINDOW    = "Window2"
DEFAULT_HARMONICS = 5


def fetch_nvgate_windows():
    """Return list of NVGate window names, or [] if NVGate is not reachable."""
    try:
        with connect() as client:
            cmd = GetWindowNameList()
            client.send_command(cmd)
            return cmd.value or []
    except Exception:
        return []


def run_thd_calculation(window_name, num_harmonics, log):
    """
    Full THD calculation pipeline.
    log(msg) is called with each progress line.
    Returns (result_dict, None) on success or (None, error_string) on failure.
    """
    try:
        with connect() as client:

            # 1. Trace info
            log(f"Reading spectrum from '{window_name}'...")
            cmd = GetWindowTraces(window_name=window_name)
            client.send_command(cmd)
            if not cmd.list_result_format_and_description:
                return None, f"'{window_name}' is empty — open a spectrum there first."

            trace   = cmd.list_result_format_and_description[0]
            module  = trace.module_id
            process = trace.process_id
            channel = trace.channel_number
            ref_ch  = trace.reference_channel_number or 0

            # 2. Download spectrum
            log("Downloading spectrum data...")
            try:
                client.send_command(DeleteAllResultEx())
            except pynvdrive.NVDriveCommandError:
                pass
            client.send_command(
                CreateResultEx(module=module, process=process,
                               channel=channel, reference_channel_number=ref_ch)
            )
            cmd_r = GetResultEx(module=module, process=process,
                                channel=channel, reference_channel_number=ref_ch)
            client.send_command(cmd_r)
            result = cmd_r.result

            # 3. Amplitude list + frequency axis
            amplitudes = get_amplitudes(result)
            if amplitudes is None:
                return None, f"'{window_name}' does not contain a spectrum (no RVEC/CVEC)."
            if result.xreg is None:
                return None, "No frequency axis found — only linear FFT spectra are supported."

            freq_start = result.xreg.start
            freq_stop  = result.xreg.stop
            log(f"{len(amplitudes)} bins,  {freq_start:.2f} - {freq_stop:.2f} Hz")

            # 4. Fundamental
            _, f0, h1 = find_fundamental(amplitudes, freq_start, freq_stop)
            log(f"Fundamental: f0 = {f0:.4f} Hz   H1 = {h1:.6g}")

            max_harm_freq = num_harmonics * f0
            if max_harm_freq > freq_stop:
                log(f"Warning: H{num_harmonics} ({max_harm_freq:.1f} Hz) "
                    f"exceeds spectrum range ({freq_stop:.1f} Hz).")

            # 5. Display zone
            display_zone = ''
            try:
                cmd_dz = GetWindowDisplayZone(window_name=window_name)
                client.send_command(cmd_dz)
                if cmd_dz.list_display_zone_type:
                    display_zone = cmd_dz.list_display_zone_type[0]
            except pynvdrive.NVDriveCommandError:
                pass

            # 6. Harmonic marker
            log(f"Adding Harmonic marker ({num_harmonics} harmonics)...")
            try:
                client.send_command(
                    AddHarmonicMarker(
                        window_name=window_name,
                        start_position=f0,
                        num_harmonics=num_harmonics,
                        display_zone=display_zone,
                    )
                )
                log("  Marker added.")
            except pynvdrive.NVDriveCommandError as e:
                log(f"  Warning: marker not added ({e.error}). Continuing.")

            # 7. Harmonic amplitudes H2 ... Hn
            log("Extracting harmonic amplitudes...")
            overtones = []
            for n in range(2, num_harmonics + 1):
                amp = amplitude_at(amplitudes, freq_start, freq_stop, n * f0)
                overtones.append(amp)
                log(f"  H{n} = {amp:.6g}  ({n * f0:.2f} Hz)")

            # 8. THD
            thd = compute_thd(h1, overtones)
            log(f"\n  * THD = {thd:.2f} %\n")

            # 9. Inject scalar + open window
            log("Sending result to NVGate...")
            channel_id = inject_scalar_percent(client, thd, name="THD")

            try:
                cmd_lay = GetCurrentLayout()
                client.send_command(cmd_lay)
                layout = cmd_lay.value or "Layout1"
            except pynvdrive.NVDriveCommandError:
                layout = "Layout1"

            client.send_command(
                AddWindowResult(
                    layout=layout, window="THD",
                    module=498, process=182, channel=channel_id,
                    existing_result_is_an_error='0',
                )
            )
            log(f"Window 'THD' opened in layout '{layout}'.")

            # 10. Gather metadata (best-effort; no failure if any sub-command fails)
            log("Collecting metadata...")
            meta = _gather_metadata(client, result, trace, layout, window_name)

            return {
                'thd':          thd,
                'f0':           f0,
                'h1':           h1,
                'overtones':    overtones,
                'amplitudes':   amplitudes,
                'freq_start':   freq_start,
                'freq_stop':    freq_stop,
                'num_harmonics': num_harmonics,
                'window_name':  window_name,
                'meta':         meta,
            }, None

    except pynvdrive.NVDriveCommandError as e:
        return None, f"NVGate error: {e.error}"
    except ConnectionError as e:
        return None, str(e)
    except Exception as e:
        return None, f"Unexpected error: {e}"


# ── Excel export ─────────────────────────────────────────────────────────────

def export_to_excel(filepath, data):
    """
    Write a 3-sheet .xlsx file:
      Sheet "Metadata"     — NVGate context, signal info, FFT parameters
      Sheet "Spectrum"     — full freq/amplitude table + scatter chart
      Sheet "THD Analysis" — harmonic table, formula, result
    """
    from openpyxl import Workbook
    from openpyxl.chart import ScatterChart, Reference, Series
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

    wb = Workbook()

    # ── style constants ───────────────────────────────────────────────────────
    HEADER_FONT   = Font(bold=True)
    TITLE_FONT    = Font(bold=True, size=13)
    SECTION_FONT  = Font(bold=True, color="1F497D")
    GREY_FILL     = PatternFill("solid", fgColor="DAEEF3")
    THIN          = Side(style="thin", color="B0B0B0")
    BOX           = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    amplitudes  = data['amplitudes']
    freq_start  = data['freq_start']
    freq_stop   = data['freq_stop']
    f0          = data['f0']
    h1          = data['h1']
    overtones   = data['overtones']
    num_harm    = data['num_harmonics']
    thd         = data['thd']
    window_name = data['window_name']
    meta        = data.get('meta', {})
    n_bins      = len(amplitudes)
    df          = (freq_stop - freq_start) / max(n_bins - 1, 1)

    # ─────────────────────────────────────────────────────────────────────────
    # Sheet 1: Metadata  (first sheet — use the default active sheet)
    # ─────────────────────────────────────────────────────────────────────────
    ws_meta = wb.active
    ws_meta.title = "Metadata"
    ws_meta.column_dimensions['A'].width = 32
    ws_meta.column_dimensions['B'].width = 48

    def _meta_title(ws, row, text):
        c = ws.cell(row=row, column=1, value=text)
        c.font = TITLE_FONT
        return row + 1

    def _meta_section(ws, row, text):
        c = ws.cell(row=row, column=1, value=text)
        c.font = SECTION_FONT
        c.fill = GREY_FILL
        ws.cell(row=row, column=2).fill = GREY_FILL
        return row + 1

    def _meta_row(ws, row, label, value):
        ws.cell(row=row, column=1, value=label).font = HEADER_FONT
        ws.cell(row=row, column=2, value=value)
        return row + 1

    r = 1
    r = _meta_title(ws_meta, r, "THD Analysis — Measurement Metadata")
    r += 1   # blank line

    # NVGate & hardware
    r = _meta_section(ws_meta, r, "NVGate Application")
    for label in ('NVGate version', 'Build number', 'Hardware type', 'Serial number'):
        v = meta.get(label, '')
        if v:
            r = _meta_row(ws_meta, r, label, v)
    r += 1

    # Project & layout context
    r = _meta_section(ws_meta, r, "Project & Session Context")
    for label in ('Project', 'Database path', 'Layout', 'NVGate window'):
        r = _meta_row(ws_meta, r, label, meta.get(label, ''))
    r += 1

    # Signal & analysis
    r = _meta_section(ws_meta, r, "Signal & Analysis")
    for label in ('Signal (channel)', 'Analysis', 'Result name',
                  'Physical quantity', 'Unit', 'Magnitude type',
                  'Result type', 'Measurement status', 'Complex data'):
        v = meta.get(label, '')
        if v:
            r = _meta_row(ws_meta, r, label, v)
    r += 1

    # Measurement timing & levels
    r = _meta_section(ws_meta, r, "Measurement Info")
    for label in ('Measurement date', 'Global RMS level',
                  'A-weighted global level', 'Tacho speed at measurement'):
        v = meta.get(label, '')
        if v:
            r = _meta_row(ws_meta, r, label, v)
    r += 1

    # FFT parameters (computed from spectrum data)
    r = _meta_section(ws_meta, r, "FFT Parameters")
    r = _meta_row(ws_meta, r, "Spectrum type",        meta.get('Spectrum type', 'Linear (FFT)'))
    r = _meta_row(ws_meta, r, "Spectral window",      meta.get('Spectral window', ''))
    r = _meta_row(ws_meta, r, "Frequency min (Hz)",   round(freq_start, 4))
    r = _meta_row(ws_meta, r, "Frequency max (Hz)",   round(freq_stop,  4))
    r = _meta_row(ws_meta, r, "X-axis unit",          meta.get('X-axis unit', 'Hz'))
    r = _meta_row(ws_meta, r, "Number of lines",      n_bins)
    r = _meta_row(ws_meta, r, "Frequency resolution", round(df, 6))
    ws_meta.cell(row=r - 1, column=2).value   # keep format consistent
    r = _meta_row(ws_meta, r, "Peak amplitude",       h1)
    r = _meta_row(ws_meta, r, "Frequency at peak (Hz)", round(f0, 4))
    r += 1

    # THD calculation parameters
    r = _meta_section(ws_meta, r, "THD Calculation")
    r = _meta_row(ws_meta, r, "Number of harmonics", num_harm)
    r = _meta_row(ws_meta, r, "THD result (%)",       round(thd, 4))
    r += 1

    # Export provenance
    r = _meta_section(ws_meta, r, "Export")
    r = _meta_row(ws_meta, r, "Exported by", meta.get('Exported by', 'thd-ui'))
    r = _meta_row(ws_meta, r, "Export date", meta.get('Export date',
        datetime.datetime.now().strftime('%Y-%m-%d  %H:%M:%S')))

    # ─────────────────────────────────────────────────────────────────────────
    # Sheet 2: Spectrum
    # ─────────────────────────────────────────────────────────────────────────
    ws1 = wb.create_sheet("Spectrum")
    ws1.column_dimensions['A'].width = 18
    ws1.column_dimensions['B'].width = 18
    ws1.column_dimensions['C'].width = 20
    ws1.column_dimensions['D'].width = 20

    for col, label in [
        (1, "Frequency (Hz)"), (2, "Amplitude"),
        (3, "Harmonic Freq (Hz)"), (4, "Harmonic Amplitude"),
    ]:
        cell = ws1.cell(row=1, column=col, value=label)
        cell.font = HEADER_FONT
        cell.fill = GREY_FILL

    # Full spectrum
    for i, amp in enumerate(amplitudes):
        ws1.cell(row=i + 2, column=1, value=round(freq_start + i * df, 6))
        ws1.cell(row=i + 2, column=2, value=amp)

    # Harmonic positions H1...Hn (columns C/D — used by the chart)
    harmonic_freqs = [f0 * k for k in range(1, num_harm + 1)]
    harmonic_amps  = [h1] + overtones
    for i, (hf, ha) in enumerate(zip(harmonic_freqs, harmonic_amps)):
        ws1.cell(row=i + 2, column=3, value=round(hf, 4))
        ws1.cell(row=i + 2, column=4, value=ha)

    # Scatter chart
    unit_lbl = meta.get('Unit', '')
    chart = ScatterChart()
    chart.title        = f"FFT Spectrum  -  {window_name}"
    chart.style        = 10
    chart.width        = 22
    chart.height       = 14
    chart.y_axis.title = f"Amplitude  ({unit_lbl})" if unit_lbl else "Amplitude"
    chart.x_axis.title = "Frequency  (Hz)"

    x_spec = Reference(ws1, min_col=1, min_row=2, max_row=n_bins + 1)
    y_spec = Reference(ws1, min_col=2, min_row=2, max_row=n_bins + 1)
    s_spec = Series(y_spec, x_spec, title="Spectrum")
    s_spec.marker.symbol = "none"
    s_spec.graphicalProperties.line.width = 12700
    chart.series.append(s_spec)

    x_harm = Reference(ws1, min_col=3, min_row=2, max_row=num_harm + 1)
    y_harm = Reference(ws1, min_col=4, min_row=2, max_row=num_harm + 1)
    s_harm = Series(y_harm, x_harm, title="Harmonics")
    s_harm.marker.symbol = "diamond"
    s_harm.marker.size   = 8
    s_harm.graphicalProperties.line.noFill = True
    chart.series.append(s_harm)

    ws1.add_chart(chart, "F2")

    # ─────────────────────────────────────────────────────────────────────────
    # Sheet 3: THD Analysis
    # ─────────────────────────────────────────────────────────────────────────
    ws2 = wb.create_sheet("THD Analysis")
    ws2.column_dimensions['A'].width = 26
    ws2.column_dimensions['B'].width = 22
    ws2.column_dimensions['C'].width = 18

    row = 1
    ws2.cell(row=row, column=1, value="THD Analysis Report").font = TITLE_FONT
    row += 2

    for label, value in [
        ("NVGate window",       window_name),
        ("Number of harmonics", num_harm),
        ("Date",                datetime.datetime.now().strftime("%Y-%m-%d  %H:%M")),
    ]:
        ws2.cell(row=row, column=1, value=label).font = HEADER_FONT
        ws2.cell(row=row, column=2, value=value)
        row += 1
    row += 1

    # Harmonics table
    ws2.cell(row=row, column=1, value="Harmonics").font = SECTION_FONT
    row += 1

    for col, label in [(1, "Harmonic"), (2, "Frequency (Hz)"), (3, "Amplitude")]:
        cell = ws2.cell(row=row, column=col, value=label)
        cell.font   = HEADER_FONT
        cell.fill   = GREY_FILL
        cell.border = BOX
    row += 1

    ws2.cell(row=row, column=1, value="H1 - Fundamental").border = BOX
    ws2.cell(row=row, column=2, value=round(f0, 4)).border       = BOX
    ws2.cell(row=row, column=3, value=h1).border                  = BOX
    row += 1

    for k, amp in enumerate(overtones, start=2):
        ws2.cell(row=row, column=1, value=f"H{k}").border          = BOX
        ws2.cell(row=row, column=2, value=round(f0 * k, 4)).border = BOX
        ws2.cell(row=row, column=3, value=amp).border               = BOX
        row += 1
    row += 1

    # THD formula
    ws2.cell(row=row, column=1, value="THD Formula").font = SECTION_FONT
    row += 1

    harm_terms   = " + ".join(f"H{k}^2" for k in range(2, num_harm + 1))
    formula_text = f"THD = sqrt({harm_terms}) / H1 x 100"
    ws2.cell(row=row, column=1, value=formula_text)
    ws2.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    row += 1

    sub_terms = " + ".join(f"{amp:.4g}^2" for amp in overtones)
    sub_text  = f"THD = sqrt({sub_terms}) / {h1:.4g} x 100"
    ws2.cell(row=row, column=1, value=sub_text)
    ws2.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    row += 2

    result_font = Font(bold=True, size=12)
    ws2.cell(row=row, column=1, value="THD Result").font  = result_font
    ws2.cell(row=row, column=2, value=round(thd, 4)).font = result_font
    ws2.cell(row=row, column=3, value="%").font            = result_font

    # Make Metadata the active (first-visible) sheet
    wb.active = ws_meta

    wb.save(filepath)


# ── Tkinter application ─────────────────────────────────────────────────────

class THDApp(tk.Tk):

    def __init__(self):
        super().__init__()
        self.title("THD Calculator")
        self.resizable(False, False)
        self._last_calc = None
        self._build_ui()
        self.after(150, self._refresh_windows)

    # ── UI construction ─────────────────────────────────────────────────────

    def _build_ui(self):
        PAD = 14
        f = ttk.Frame(self, padding=PAD)
        f.grid(row=0, column=0, sticky="nsew")

        ttk.Label(f, text="NVGate window:").grid(
            row=0, column=0, sticky="w", pady=(0, 6))

        self._window_var = tk.StringVar(value=DEFAULT_WINDOW)
        self._window_combo = ttk.Combobox(
            f, textvariable=self._window_var, width=24, state="readonly")
        self._window_combo.grid(
            row=0, column=1, sticky="ew", padx=(8, 4), pady=(0, 6))

        self._refresh_btn = ttk.Button(f, text="↺", width=3,
                                       command=self._refresh_windows)
        self._refresh_btn.grid(row=0, column=2, pady=(0, 6))

        ttk.Label(f, text="Number of harmonics:").grid(
            row=1, column=0, sticky="w", pady=(0, 12))

        vcmd = (self.register(self._validate_positive_int), '%P')
        self._harmonics_var = tk.StringVar(value=str(DEFAULT_HARMONICS))
        ttk.Entry(f, textvariable=self._harmonics_var, width=6,
                  validate='key', validatecommand=vcmd).grid(
            row=1, column=1, sticky="w", padx=(8, 0), pady=(0, 12))

        self._run_btn = ttk.Button(f, text="▶   Run THD", width=32,
                                   command=self._on_run, state="disabled")
        self._run_btn.grid(row=2, column=0, columnspan=3, pady=(0, 4))

        self._export_btn = ttk.Button(f, text="Export to Excel", width=32,
                                      command=self._on_export, state="disabled")
        self._export_btn.grid(row=3, column=0, columnspan=3, pady=(0, 10))

        ttk.Separator(f, orient="horizontal").grid(
            row=4, column=0, columnspan=3, sticky="ew", pady=(0, 8))

        self._log_box = tk.Text(
            f, height=13, width=54, state="disabled",
            font=("Courier", 9), wrap="word",
            bg="#f5f5f5", relief="flat", cursor="arrow")
        self._log_box.grid(row=5, column=0, columnspan=3, sticky="nsew")

        sb = ttk.Scrollbar(f, command=self._log_box.yview)
        sb.grid(row=5, column=3, sticky="ns")
        self._log_box.configure(yscrollcommand=sb.set)

        f.columnconfigure(1, weight=1)

    # ── Validation ──────────────────────────────────────────────────────────

    @staticmethod
    def _validate_positive_int(value):
        return value == "" or value.isdigit()

    # ── Window list ─────────────────────────────────────────────────────────

    def _refresh_windows(self):
        self._refresh_btn.configure(state="disabled")
        self._run_btn.configure(state="disabled")
        self._log("Connecting to NVGate...")

        def worker():
            windows = fetch_nvgate_windows()
            self.after(0, lambda: self._on_windows_loaded(windows))

        threading.Thread(target=worker, daemon=True).start()

    def _on_windows_loaded(self, windows):
        self._refresh_btn.configure(state="normal")
        if not windows:
            self._log("Could not reach NVGate. "
                      "Start NVGate in connected mode, then click ↺.")
            return

        self._window_combo.configure(values=windows)
        current = self._window_var.get()
        if current not in windows:
            self._window_var.set(
                DEFAULT_WINDOW if DEFAULT_WINDOW in windows else windows[0])

        self._run_btn.configure(state="normal")
        self._log(f"Ready — {len(windows)} window(s): {', '.join(windows)}")

    # ── Run ─────────────────────────────────────────────────────────────────

    def _on_run(self):
        window = self._window_var.get()
        if not window:
            self._log("Please select a window first.")
            return

        raw = self._harmonics_var.get().strip()
        if not raw or not raw.isdigit() or int(raw) < 2:
            self._log("Number of harmonics must be a whole number >= 2.")
            return
        num_harmonics = int(raw)

        self._run_btn.configure(state="disabled")
        self._refresh_btn.configure(state="disabled")
        self._export_btn.configure(state="disabled")
        self._last_calc = None
        self._clear_log()
        self._log(f"Window: '{window}'   Harmonics: {num_harmonics}\n")

        def worker():
            calc, err = run_thd_calculation(window, num_harmonics, self._safe_log)
            self.after(0, lambda: self._on_done(calc, err))

        threading.Thread(target=worker, daemon=True).start()

    def _on_done(self, calc, error):
        self._run_btn.configure(state="normal")
        self._refresh_btn.configure(state="normal")
        if error:
            self._log(f"\nERROR: {error}")
        else:
            self._last_calc = calc
            self._export_btn.configure(state="normal")
            self._log(
                f"Done.  THD = {calc['thd']:.2f} %  —  "
                "result visible in NVGate.\n"
                "Click 'Export to Excel' to save the full report."
            )

    # ── Export ───────────────────────────────────────────────────────────────

    def _on_export(self):
        if self._last_calc is None:
            return

        window_slug  = self._last_calc['window_name'].replace(" ", "_")
        default_name = f"THD_{window_slug}.xlsx"

        filepath = filedialog.asksaveasfilename(
            title="Save Excel report",
            defaultextension=".xlsx",
            initialfile=default_name,
            filetypes=[("Excel workbook", "*.xlsx"), ("All files", "*.*")],
        )
        if not filepath:
            return

        self._export_btn.configure(state="disabled")
        self._log("Writing Excel file...")

        data = self._last_calc

        def worker():
            try:
                export_to_excel(filepath, data)
                self.after(0, lambda: self._on_export_done(filepath, None))
            except Exception as e:
                self.after(0, lambda: self._on_export_done(filepath, str(e)))

        threading.Thread(target=worker, daemon=True).start()

    def _on_export_done(self, filepath, error):
        self._export_btn.configure(state="normal")
        if error:
            self._log(f"Export failed: {error}")
        else:
            self._log(f"Saved: {filepath}")

    # ── Log helpers ─────────────────────────────────────────────────────────

    def _log(self, msg):
        self._log_box.configure(state="normal")
        self._log_box.insert("end", msg + "\n")
        self._log_box.see("end")
        self._log_box.configure(state="disabled")

    def _safe_log(self, msg):
        self.after(0, lambda: self._log(msg))

    def _clear_log(self):
        self._log_box.configure(state="normal")
        self._log_box.delete("1.0", "end")
        self._log_box.configure(state="disabled")


# ── Entry point ─────────────────────────────────────────────────────────────

if __name__ == "__main__":
    THDApp().mainloop()
