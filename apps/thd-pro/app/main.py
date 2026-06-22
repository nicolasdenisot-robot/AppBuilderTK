"""
thd-pro — THD Analyzer with a modern PySide6/Qt UI.

Results (THD %, THD dB, H1 frequency, H1 amplitude) are displayed in the
application window. A harmonic marker is placed on the NVGate spectrum for
visual reference. "Export to Excel" saves a full 3-sheet report.
"""

import sys
import os
import math
import datetime
from pathlib import Path

# ── Path setup: works inside workspace and in an exported bundle ─────────────
_HERE = Path(__file__).resolve()
for _candidate in (_HERE.parents[3] / "shared", _HERE.parents[1]):
    if (_candidate / "oros_helpers").is_dir() and str(_candidate) not in sys.path:
        sys.path.insert(0, str(_candidate))

import oros_helpers  # noqa: F401  — side-effect: calls ensure_pynvdrive()
from oros_helpers import connect

import pynvdrive
from pynvdrive.command import CommandAction
from pynvdrive.commands.marker import GetWindowTraces, GetWindowDisplayZone
from pynvdrive.commands.graphics.getwindownamelist import GetWindowNameList
from pynvdrive.commands.signalresult import (
    CreateResultEx, GetResultEx, DeleteAllResultEx,
)
from pynvdrive.commands.project.getcurrentprojectname import GetCurrentProjectName
from pynvdrive.commands.project.getdatabasepath import GetDatabasePath
from pynvdrive.commands.analyzers.getnvgateinfo import GetNVGateInfo

from PySide6.QtCore import Qt, Signal, QThread
from PySide6.QtGui import QPixmap
from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QGridLayout, QLabel, QComboBox, QPushButton, QLineEdit,
    QFileDialog, QFrame, QSizePolicy,
)

_LOGO_BRAND   = "oros"   # folder under shared/assets/logos/
_LOGO_VARIANT = "color"  # "color" for light themes · "white" for dark themes

def _logo_path() -> str:
    """Return the logo path — works in development (shared/assets) and as exe (sys._MEIPASS)."""
    fname = f"logo_{_LOGO_VARIANT}.png"
    if hasattr(sys, '_MEIPASS'):
        return os.path.join(sys._MEIPASS, fname)
    _here = Path(__file__).resolve()
    for candidate in (
        _here.parents[3] / "shared" / "assets" / "logos" / _LOGO_BRAND,
        _here.parent,  # local override next to main.py
    ):
        p = candidate / fname
        if p.exists():
            return str(p)
    return ""

# ── Lookup tables for coded result-block values ───────────────────────────────
_RESULT_TYPE = {1: "Scalar", 2: "Spectrum (regular vector)",
                3: "Reference vector", 8: "Waterfall"}
_MEASURE_STATUS = {0: "No measurement", 1: "In progress", 2: "Finished"}
_N_OCTAVE = {0: "Linear (FFT)", 1: "1/1 octave", 2: "1/3 octave",
             3: "1/12 octave", 4: "1/24 octave"}
_WEIGHTING_WINDOW = {
    0: "Rectangular (none)", 1: "Hann", 2: "Flat-top",
    3: "Hamming", 4: "Blackman", 5: "Exponential",
    6: "Force", 7: "Exponential + Force", 8: "Kaiser-Bessel",
}

# ── Application stylesheet (light theme) ─────────────────────────────────────
_STYLESHEET = """
QWidget {
    font-family: "Segoe UI", Arial, sans-serif;
    font-size: 13px;
    background-color: #f4f6f8;
    color: #1a1a2e;
}
QFrame#ConfigFrame {
    background-color: #ffffff;
    border: 1px solid #e0e4ea;
    border-radius: 8px;
}
QFrame#ResultCard {
    background-color: #ffffff;
    border: 1px solid #e0e4ea;
    border-radius: 8px;
}
QLabel#AppTitle {
    font-size: 20px;
    font-weight: 700;
    color: #1a1a2e;
    background: transparent;
}
QLabel#AppSubtitle {
    font-size: 11px;
    color: #6b7280;
    background: transparent;
}
QLabel#CardLabel {
    font-size: 11px;
    color: #6b7280;
    font-weight: 400;
    background: transparent;
}
QLabel#CardValue {
    font-size: 28px;
    font-weight: 700;
    color: #9ca3af;
    background: transparent;
}
QLabel#CardUnit {
    font-size: 15px;
    font-weight: 400;
    color: #9ca3af;
    background: transparent;
    padding-bottom: 3px;
}
QLabel#ConfigLabel {
    color: #374151;
    background: transparent;
}
QPushButton#RunButton {
    background-color: #1565c0;
    color: white;
    border: none;
    border-radius: 6px;
    font-size: 14px;
    font-weight: 600;
    min-height: 42px;
    padding: 0px 24px;
}
QPushButton#RunButton:hover {
    background-color: #1976d2;
}
QPushButton#RunButton:pressed {
    background-color: #0d47a1;
}
QPushButton#RunButton:disabled {
    background-color: #b0bec5;
    color: #eceff1;
}
QPushButton#ExportButton {
    background-color: #ffffff;
    color: #1565c0;
    border: 2px solid #1565c0;
    border-radius: 6px;
    font-size: 14px;
    font-weight: 600;
    min-height: 42px;
    padding: 0px 24px;
}
QPushButton#ExportButton:hover {
    background-color: #e3f2fd;
}
QPushButton#ExportButton:pressed {
    background-color: #bbdefb;
}
QPushButton#ExportButton:disabled {
    color: #b0bec5;
    border-color: #d1d5db;
    background-color: #f9fafb;
}
QComboBox {
    background-color: #ffffff;
    border: 1px solid #d1d5db;
    border-radius: 5px;
    padding: 5px 8px;
    color: #1a1a2e;
    min-width: 160px;
    selection-background-color: #1565c0;
}
QComboBox:focus {
    border: 1px solid #1565c0;
}
QComboBox::drop-down {
    border: none;
    padding-right: 4px;
}
QComboBox QAbstractItemView {
    background-color: #ffffff;
    border: 1px solid #d1d5db;
    selection-background-color: #e3f2fd;
    selection-color: #1a1a2e;
}
QPushButton#RefreshButton {
    background-color: #f3f4f6;
    border: 1px solid #d1d5db;
    border-radius: 5px;
    font-size: 15px;
    color: #374151;
    min-width: 32px;
    max-width: 32px;
    min-height: 30px;
    max-height: 30px;
    padding: 0px;
}
QPushButton#RefreshButton:hover {
    background-color: #e5e7eb;
    border-color: #9ca3af;
}
QPushButton#StepButton {
    background-color: #e8f0fe;
    color: #1565c0;
    border: 1.5px solid #aac4f0;
    border-radius: 14px;
    font-size: 20px;
    font-weight: 600;
    min-width: 28px;
    max-width: 28px;
    min-height: 28px;
    max-height: 28px;
    padding: 0px;
}
QPushButton#StepButton:hover {
    background-color: #c5d9f8;
    border-color: #1565c0;
}
QPushButton#StepButton:pressed {
    background-color: #90b8f0;
}
QPushButton#StepButton:disabled {
    background-color: #f3f4f6;
    color: #d1d5db;
    border-color: #e5e7eb;
}
QLineEdit#StepDisplay {
    font-size: 16px;
    font-weight: 600;
    color: #1a1a2e;
    background: #ffffff;
    border: 1.5px solid #d1d5db;
    border-radius: 5px;
    padding: 2px 4px;
    min-width: 38px;
    max-width: 38px;
}
QLineEdit#StepDisplay:focus {
    border-color: #1565c0;
}
"""

# Status panel inline styles (set dynamically per state)
_STATUS_STYLES = {
    "info":    ("#e8f4fd", "#1a3a5c", "#b3d9f7", "#1565c0"),   # bg, text, border, dot
    "success": ("#e8f5e9", "#1b5e20", "#a5d6a7", "#2e7d32"),
    "error":   ("#fdecea", "#7f1d1d", "#f5a5a0", "#c62828"),
    "working": ("#fff8e1", "#4a3200", "#ffe082", "#f57f17"),
}

_SUP_TRANS = str.maketrans("0123456789+-", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻")

def _fmt_sci(v: float, digits: int = 3) -> str:
    """Format a float as clean scientific notation: 5.183 × 10⁻⁴"""
    if v == 0:
        return "0"
    mantissa, exp = f"{v:.{digits}e}".split("e")
    return f"{mantissa} × 10{str(int(exp)).translate(_SUP_TRANS)}"


# ── Custom NVDrive command ────────────────────────────────────────────────────

class AddHarmonicMarker(CommandAction):
    """Add a harmonic marker on a display zone (NVDrive manual §3.2.9)."""
    def __init__(self, window_name, start_position, num_harmonics=5, display_zone=''):
        super().__init__(
            name=None,
            parameters=[window_name, str(start_position), str(num_harmonics),
                        str(display_zone)],
        )


# ── THD mathematics ───────────────────────────────────────────────────────────

def get_amplitudes(result):
    if result.rvec is not None:
        return list(result.rvec.data)
    if result.cvec is not None:
        return [math.sqrt(d[0] ** 2 + d[1] ** 2) for d in result.cvec.data]
    return None


def find_fundamental(amplitudes, freq_start, freq_stop):
    idx = max(range(len(amplitudes)), key=lambda i: amplitudes[i])
    n = len(amplitudes)
    freq = freq_start + idx * (freq_stop - freq_start) / max(n - 1, 1)
    return idx, freq, amplitudes[idx]


def amplitude_at(amplitudes, freq_start, freq_stop, target_hz):
    n = len(amplitudes)
    df = (freq_stop - freq_start) / max(n - 1, 1)
    idx = int(round((target_hz - freq_start) / df))
    idx = max(0, min(idx, n - 1))
    return amplitudes[idx]


def compute_thd(h1, overtones):
    if abs(h1) == 0:
        raise ValueError("Fundamental amplitude is zero — cannot compute THD.")
    return math.sqrt(sum(h ** 2 for h in overtones)) / abs(h1) * 100.0


# ── Metadata helpers ──────────────────────────────────────────────────────────

def _safe_get(obj, attr, default=''):
    val = getattr(obj, attr, None)
    return val if val is not None else default


def _decode(code, lookup, default_prefix='code'):
    if code is None:
        return ''
    return lookup.get(code, f'{default_prefix} {code}')


def _gather_metadata(client, result, trace, window_name):
    meta = {}

    try:
        cmd = GetNVGateInfo()
        client.send_command(cmd)
        meta['NVGate version'] = f"{cmd.nvgate_major_version}.{cmd.nvgate_minor_version}"
        meta['Build number']   = _safe_get(cmd, 'full_build_number')
        meta['Hardware type']  = _safe_get(cmd, 'hard_type')
        meta['Serial number']  = _safe_get(cmd, 'serial_number')
    except Exception:
        meta['NVGate version'] = meta['Build number'] = ''
        meta['Hardware type']  = meta['Serial number'] = ''

    try:
        cmd = GetCurrentProjectName()
        client.send_command(cmd)
        meta['Project'] = cmd.value or '(not saved / live session)'
    except Exception:
        meta['Project'] = ''

    try:
        cmd = GetDatabasePath()
        client.send_command(cmd)
        meta['Database path'] = cmd.value or ''
    except Exception:
        meta['Database path'] = ''

    meta['NVGate window']     = window_name
    meta['Signal (channel)']  = trace.input_name or ''
    try:
        meta['Analysis'] = trace.analysis_fullname
    except Exception:
        meta['Analysis'] = f"module {trace.module_id}, process {trace.process_id}"

    meta['Result name']       = _safe_get(result.head, 'name')
    meta['Physical quantity'] = _safe_get(result.head, 'physical_quantity')
    meta['Unit']              = (_safe_get(result.head, 'unit_label')
                                 or _safe_get(result.head, 'unit_name'))
    meta['Magnitude type']    = _safe_get(result.head, 'magnitude_key')

    rt   = _safe_get(result.head, 'result_type')
    ms   = _safe_get(result.head, 'measure_status')
    ww   = _safe_get(result.head, 'weighting_window')
    is_c = _safe_get(result.head, 'is_complex')

    meta['Result type']        = (_decode(rt, _RESULT_TYPE,    'type')
                                  + (f'  ({rt})' if rt != '' else ''))
    meta['Measurement status'] = (_decode(ms, _MEASURE_STATUS, 'status')
                                  + (f'  ({ms})' if ms != '' else ''))
    meta['Spectral window']    = (_decode(ww, _WEIGHTING_WINDOW, 'code')
                                  + (f'  ({ww})' if ww != '' else ''))
    meta['Complex data']       = 'Yes' if is_c else 'No'

    if result.info is not None:
        d = _safe_get(result.info, 'date')
        try:
            meta['Measurement date'] = d.strftime('%Y-%m-%d  %H:%M:%S') if d else ''
        except Exception:
            meta['Measurement date'] = str(d)
        gl  = _safe_get(result.info, 'global_level', 0.0)
        wgl = _safe_get(result.info, 'weighted_global_level', 0.0)
        tch = _safe_get(result.info, 'tacho', 0.0)
        unit_lbl = meta['Unit'] or ''
        meta['Global RMS level']           = f"{gl:.6g}  {unit_lbl}" if gl else ''
        meta['A-weighted global level']    = f"{wgl:.6g}  {unit_lbl}" if wgl else ''
        meta['Tacho speed at measurement'] = f"{tch:.4g} RPM" if tch else ''
    else:
        meta['Measurement date'] = meta['Global RMS level'] = ''
        meta['A-weighted global level'] = meta['Tacho speed at measurement'] = ''

    if result.xreg is not None:
        no = _safe_get(result.xreg, 'n_octave')
        meta['Spectrum type'] = _decode(no, _N_OCTAVE, 'type') if no != '' else ''
        meta['X-axis unit']   = (_safe_get(result.xreg, 'unit_label')
                                 or _safe_get(result.xreg, 'unit_name'))
    else:
        meta['Spectrum type'] = meta['X-axis unit'] = ''

    meta['Exported by'] = 'THD Analyzer  (thd-pro)'
    meta['Export date'] = datetime.datetime.now().strftime('%Y-%m-%d  %H:%M:%S')
    return meta


# ── NVGate helpers ────────────────────────────────────────────────────────────

DEFAULT_WINDOW    = "Window2"
DEFAULT_HARMONICS = 5


def fetch_nvgate_windows():
    try:
        with connect() as client:
            cmd = GetWindowNameList()
            client.send_command(cmd)
            return cmd.value or []
    except Exception:
        return []


def run_thd_calculation(window_name, num_harmonics, log):
    try:
        with connect() as client:

            log(f"Reading spectrum from '{window_name}'…")
            cmd = GetWindowTraces(window_name=window_name)
            client.send_command(cmd)
            if not cmd.list_result_format_and_description:
                return None, f"'{window_name}' is empty — open a spectrum there first."

            trace   = cmd.list_result_format_and_description[0]
            module  = trace.module_id
            process = trace.process_id
            channel = trace.channel_number
            ref_ch  = trace.reference_channel_number or 0

            log("Downloading spectrum data…")
            try:
                client.send_command(DeleteAllResultEx())
            except pynvdrive.NVDriveCommandError:
                pass
            client.send_command(
                CreateResultEx(module=module, process=process,
                               channel=channel, reference_channel_number=ref_ch)
            )
            cmd_r = GetResultEx(module=module, process=process,
                                channel=channel, reference_channel_number=ref_ch)
            client.send_command(cmd_r)
            result = cmd_r.result

            amplitudes = get_amplitudes(result)
            if amplitudes is None:
                return None, f"'{window_name}' does not contain a spectrum (no RVEC/CVEC)."
            if result.xreg is None:
                return None, "No frequency axis found — only linear FFT spectra are supported."

            freq_start = result.xreg.start
            freq_stop  = result.xreg.stop
            log(f"{len(amplitudes)} bins,  {freq_start:.2f} – {freq_stop:.2f} Hz")

            _, f0, h1 = find_fundamental(amplitudes, freq_start, freq_stop)
            log(f"Fundamental: f₀ = {f0:.4f} Hz   H1 = {h1:.6g}")

            max_harm_freq = num_harmonics * f0
            if max_harm_freq > freq_stop:
                log(f"Warning: H{num_harmonics} ({max_harm_freq:.1f} Hz) "
                    f"exceeds spectrum range ({freq_stop:.1f} Hz).")

            display_zone = ''
            try:
                cmd_dz = GetWindowDisplayZone(window_name=window_name)
                client.send_command(cmd_dz)
                if cmd_dz.list_display_zone_type:
                    display_zone = cmd_dz.list_display_zone_type[0]
            except pynvdrive.NVDriveCommandError:
                pass

            log(f"Adding harmonic marker ({num_harmonics} harmonics)…")
            try:
                client.send_command(
                    AddHarmonicMarker(window_name=window_name, start_position=f0,
                                      num_harmonics=num_harmonics,
                                      display_zone=display_zone)
                )
                log("  Marker added.")
            except pynvdrive.NVDriveCommandError as e:
                log(f"  Warning: marker not added ({e.error}). Continuing.")

            log("Extracting harmonic amplitudes…")
            overtones = []
            for n in range(2, num_harmonics + 1):
                amp = amplitude_at(amplitudes, freq_start, freq_stop, n * f0)
                overtones.append(amp)
                log(f"  H{n} = {amp:.6g}  ({n * f0:.2f} Hz)")

            thd    = compute_thd(h1, overtones)
            thd_db = 20.0 * math.log10(thd / 100.0) if thd > 0 else float('-inf')
            log(f"THD = {thd:.3f} %  ({thd_db:.1f} dB)")

            log("Collecting metadata…")
            meta = _gather_metadata(client, result, trace, window_name)

            return {
                'thd':           thd,
                'thd_db':        thd_db,
                'f0':            f0,
                'h1':            h1,
                'overtones':     overtones,
                'amplitudes':    amplitudes,
                'freq_start':    freq_start,
                'freq_stop':     freq_stop,
                'num_harmonics': num_harmonics,
                'window_name':   window_name,
                'meta':          meta,
            }, None

    except pynvdrive.NVDriveCommandError as e:
        return None, f"NVGate error: {e.error}"
    except ConnectionError as e:
        return None, str(e)
    except Exception as e:
        return None, f"Unexpected error: {e}"


# ── Excel export ──────────────────────────────────────────────────────────────

def export_to_excel(filepath, data):
    from openpyxl import Workbook
    from openpyxl.chart import ScatterChart, Reference, Series
    from openpyxl.styles import Font, PatternFill, Border, Side

    wb = Workbook()

    HEADER_FONT  = Font(bold=True)
    TITLE_FONT   = Font(bold=True, size=13)
    SECTION_FONT = Font(bold=True, color="1F497D")
    GREY_FILL    = PatternFill("solid", fgColor="DAEEF3")
    THIN         = Side(style="thin", color="B0B0B0")
    BOX          = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    amplitudes  = data['amplitudes']
    freq_start  = data['freq_start']
    freq_stop   = data['freq_stop']
    f0          = data['f0']
    h1          = data['h1']
    overtones   = data['overtones']
    num_harm    = data['num_harmonics']
    thd         = data['thd']
    thd_db      = data['thd_db']
    window_name = data['window_name']
    meta        = data.get('meta', {})
    n_bins      = len(amplitudes)
    df          = (freq_stop - freq_start) / max(n_bins - 1, 1)

    ws_meta = wb.active
    ws_meta.title = "Metadata"
    ws_meta.column_dimensions['A'].width = 32
    ws_meta.column_dimensions['B'].width = 48

    def _title(ws, row, text):
        ws.cell(row=row, column=1, value=text).font = TITLE_FONT
        return row + 1

    def _section(ws, row, text):
        c = ws.cell(row=row, column=1, value=text)
        c.font = SECTION_FONT
        c.fill = GREY_FILL
        ws.cell(row=row, column=2).fill = GREY_FILL
        return row + 1

    def _row(ws, row, label, value):
        ws.cell(row=row, column=1, value=label).font = HEADER_FONT
        ws.cell(row=row, column=2, value=value)
        return row + 1

    r = 1
    r = _title(ws_meta, r, "THD Analysis — Measurement Metadata")
    r += 1

    r = _section(ws_meta, r, "NVGate Application")
    for label in ('NVGate version', 'Build number', 'Hardware type', 'Serial number'):
        v = meta.get(label, '')
        if v:
            r = _row(ws_meta, r, label, v)
    r += 1

    r = _section(ws_meta, r, "Project & Session Context")
    for label in ('Project', 'Database path', 'NVGate window'):
        r = _row(ws_meta, r, label, meta.get(label, ''))
    r += 1

    r = _section(ws_meta, r, "Signal & Analysis")
    for label in ('Signal (channel)', 'Analysis', 'Result name',
                  'Physical quantity', 'Unit', 'Magnitude type',
                  'Result type', 'Measurement status', 'Complex data'):
        v = meta.get(label, '')
        if v:
            r = _row(ws_meta, r, label, v)
    r += 1

    r = _section(ws_meta, r, "Measurement Info")
    for label in ('Measurement date', 'Global RMS level',
                  'A-weighted global level', 'Tacho speed at measurement'):
        v = meta.get(label, '')
        if v:
            r = _row(ws_meta, r, label, v)
    r += 1

    r = _section(ws_meta, r, "FFT Parameters")
    r = _row(ws_meta, r, "Spectrum type",            meta.get('Spectrum type', 'Linear (FFT)'))
    r = _row(ws_meta, r, "Spectral window",           meta.get('Spectral window', ''))
    r = _row(ws_meta, r, "Frequency min (Hz)",        round(freq_start, 4))
    r = _row(ws_meta, r, "Frequency max (Hz)",        round(freq_stop,  4))
    r = _row(ws_meta, r, "X-axis unit",               meta.get('X-axis unit', 'Hz'))
    r = _row(ws_meta, r, "Number of lines",           n_bins)
    r = _row(ws_meta, r, "Frequency resolution (Hz)", round(df, 6))
    r = _row(ws_meta, r, "Peak amplitude",             h1)
    r = _row(ws_meta, r, "Frequency at peak (Hz)",    round(f0, 4))
    r += 1

    r = _section(ws_meta, r, "THD Calculation")
    r = _row(ws_meta, r, "Number of harmonics", num_harm)
    r = _row(ws_meta, r, "THD result (%)",       round(thd,    4))
    r = _row(ws_meta, r, "THD result (dB)",      round(thd_db, 2))
    r += 1

    r = _section(ws_meta, r, "Export")
    r = _row(ws_meta, r, "Exported by", meta.get('Exported by', 'thd-pro'))
    r = _row(ws_meta, r, "Export date", meta.get('Export date',
        datetime.datetime.now().strftime('%Y-%m-%d  %H:%M:%S')))

    ws1 = wb.create_sheet("Spectrum")
    ws1.column_dimensions['A'].width = 18
    ws1.column_dimensions['B'].width = 18
    ws1.column_dimensions['C'].width = 20
    ws1.column_dimensions['D'].width = 20

    for col, label in [(1, "Frequency (Hz)"), (2, "Amplitude"),
                       (3, "Harmonic Freq (Hz)"), (4, "Harmonic Amplitude")]:
        cell = ws1.cell(row=1, column=col, value=label)
        cell.font = HEADER_FONT
        cell.fill = GREY_FILL

    for i, amp in enumerate(amplitudes):
        ws1.cell(row=i + 2, column=1, value=round(freq_start + i * df, 6))
        ws1.cell(row=i + 2, column=2, value=amp)

    harmonic_freqs = [f0 * k for k in range(1, num_harm + 1)]
    harmonic_amps  = [h1] + overtones
    for i, (hf, ha) in enumerate(zip(harmonic_freqs, harmonic_amps)):
        ws1.cell(row=i + 2, column=3, value=round(hf, 4))
        ws1.cell(row=i + 2, column=4, value=ha)

    unit_lbl = meta.get('Unit', '')
    chart = ScatterChart()
    chart.title        = f"FFT Spectrum  —  {window_name}"
    chart.style        = 10
    chart.width        = 22
    chart.height       = 14
    chart.y_axis.title = f"Amplitude  ({unit_lbl})" if unit_lbl else "Amplitude"
    chart.x_axis.title = "Frequency  (Hz)"

    x_spec = Reference(ws1, min_col=1, min_row=2, max_row=n_bins + 1)
    y_spec = Reference(ws1, min_col=2, min_row=2, max_row=n_bins + 1)
    s_spec = Series(y_spec, x_spec, title="Spectrum")
    s_spec.marker.symbol = "none"
    s_spec.graphicalProperties.line.width = 12700
    chart.series.append(s_spec)

    x_harm = Reference(ws1, min_col=3, min_row=2, max_row=num_harm + 1)
    y_harm = Reference(ws1, min_col=4, min_row=2, max_row=num_harm + 1)
    s_harm = Series(y_harm, x_harm, title="Harmonics")
    s_harm.marker.symbol = "diamond"
    s_harm.marker.size   = 8
    s_harm.graphicalProperties.line.noFill = True
    chart.series.append(s_harm)
    ws1.add_chart(chart, "F2")

    ws2 = wb.create_sheet("THD Analysis")
    ws2.column_dimensions['A'].width = 26
    ws2.column_dimensions['B'].width = 22
    ws2.column_dimensions['C'].width = 18

    row = 1
    ws2.cell(row=row, column=1, value="THD Analysis Report").font = TITLE_FONT
    row += 2

    for label, value in [("NVGate window", window_name),
                          ("Number of harmonics", num_harm),
                          ("Date", datetime.datetime.now().strftime("%Y-%m-%d  %H:%M"))]:
        ws2.cell(row=row, column=1, value=label).font = HEADER_FONT
        ws2.cell(row=row, column=2, value=value)
        row += 1
    row += 1

    ws2.cell(row=row, column=1, value="Harmonics").font = SECTION_FONT
    row += 1

    for col, label in [(1, "Harmonic"), (2, "Frequency (Hz)"), (3, "Amplitude")]:
        cell = ws2.cell(row=row, column=col, value=label)
        cell.font = HEADER_FONT
        cell.fill = GREY_FILL
        cell.border = BOX
    row += 1

    ws2.cell(row=row, column=1, value="H1 - Fundamental").border = BOX
    ws2.cell(row=row, column=2, value=round(f0, 4)).border        = BOX
    ws2.cell(row=row, column=3, value=h1).border                   = BOX
    row += 1

    for k, amp in enumerate(overtones, start=2):
        ws2.cell(row=row, column=1, value=f"H{k}").border          = BOX
        ws2.cell(row=row, column=2, value=round(f0 * k, 4)).border = BOX
        ws2.cell(row=row, column=3, value=amp).border               = BOX
        row += 1
    row += 1

    ws2.cell(row=row, column=1, value="THD Formula").font = SECTION_FONT
    row += 1
    harm_terms   = " + ".join(f"H{k}^2" for k in range(2, num_harm + 1))
    ws2.cell(row=row, column=1, value=f"THD = sqrt({harm_terms}) / H1 × 100")
    ws2.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    row += 1
    sub_terms = " + ".join(f"{amp:.4g}^2" for amp in overtones)
    ws2.cell(row=row, column=1, value=f"THD = sqrt({sub_terms}) / {h1:.4g} × 100")
    ws2.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    row += 2

    result_font = Font(bold=True, size=12)
    ws2.cell(row=row, column=1, value="THD (%)").font     = result_font
    ws2.cell(row=row, column=2, value=round(thd, 4)).font = result_font
    ws2.cell(row=row, column=3, value="%").font            = result_font
    row += 1
    ws2.cell(row=row, column=1, value="THD (dB)").font      = result_font
    ws2.cell(row=row, column=2, value=round(thd_db, 2)).font = result_font
    ws2.cell(row=row, column=3, value="dB").font              = result_font

    wb.active = ws_meta
    wb.save(filepath)


# ── UI widgets ────────────────────────────────────────────────────────────────

class HarmonicsSelector(QWidget):
    """Modern +/− stepper with a directly editable value field."""

    def __init__(self, minimum=2, maximum=20, initial=5):
        super().__init__()
        self._min   = minimum
        self._max   = maximum
        self._value = initial

        layout = QHBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(8)

        self._btn_minus = QPushButton("−")
        self._btn_minus.setObjectName("StepButton")
        self._btn_minus.clicked.connect(self._decrement)

        self._display = QLineEdit(str(initial))
        self._display.setObjectName("StepDisplay")
        self._display.setAlignment(Qt.AlignCenter)
        self._display.setMaxLength(2)
        self._display.editingFinished.connect(self._on_edited)

        self._btn_plus = QPushButton("+")
        self._btn_plus.setObjectName("StepButton")
        self._btn_plus.clicked.connect(self._increment)

        layout.addWidget(self._btn_minus)
        layout.addWidget(self._display)
        layout.addWidget(self._btn_plus)
        layout.addStretch()

        self._refresh_buttons()

    def value(self):
        return self._value

    def _increment(self):
        if self._value < self._max:
            self._value += 1
            self._display.setText(str(self._value))
            self._refresh_buttons()

    def _decrement(self):
        if self._value > self._min:
            self._value -= 1
            self._display.setText(str(self._value))
            self._refresh_buttons()

    def _on_edited(self):
        """Validate and clamp the manually typed value."""
        try:
            v = int(self._display.text())
            self._value = max(self._min, min(self._max, v))
        except ValueError:
            pass
        self._display.setText(str(self._value))
        self._refresh_buttons()

    def _refresh_buttons(self):
        self._btn_minus.setEnabled(self._value > self._min)
        self._btn_plus.setEnabled(self._value < self._max)


class ResultCard(QFrame):
    """White card showing a metric label, large value, and inline unit."""

    def __init__(self, label, unit=""):
        super().__init__()
        self.setObjectName("ResultCard")
        self.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Preferred)
        self._unit_text = unit

        outer = QVBoxLayout(self)
        outer.setContentsMargins(18, 14, 18, 14)
        outer.setSpacing(6)

        self._label = QLabel(label)
        self._label.setObjectName("CardLabel")

        # Value and unit side by side, aligned at the bottom baseline
        value_row = QHBoxLayout()
        value_row.setSpacing(6)
        value_row.setContentsMargins(0, 0, 0, 0)

        self._value = QLabel("—")
        self._value.setObjectName("CardValue")
        self._value.setAlignment(Qt.AlignLeft | Qt.AlignBottom)

        self._unit = QLabel(unit)
        self._unit.setObjectName("CardUnit")
        self._unit.setAlignment(Qt.AlignLeft | Qt.AlignBottom)

        value_row.addWidget(self._value)
        value_row.addWidget(self._unit)
        value_row.addStretch()

        outer.addWidget(self._label)
        outer.addLayout(value_row)

    def set_ready(self, value_text, unit=""):
        self._value.setText(value_text)
        self._value.setStyleSheet(
            "font-size: 28px; font-weight: 700; color: #1565c0; background: transparent;")
        if unit:
            self._unit.setText(unit)
        self._unit.setStyleSheet(
            "font-size: 15px; font-weight: 400; color: #6b7280; background: transparent;")

    def set_empty(self):
        self._value.setText("—")
        self._value.setStyleSheet(
            "font-size: 28px; font-weight: 700; color: #9ca3af; background: transparent;")
        self._unit.setStyleSheet(
            "font-size: 15px; font-weight: 400; color: #9ca3af; background: transparent;")

    def set_unit(self, unit):
        self._unit.setText(unit)


# ── Background worker threads ─────────────────────────────────────────────────

class _WindowLoader(QThread):
    done = Signal(list)

    def run(self):
        self.done.emit(fetch_nvgate_windows())


class _ExportWorker(QThread):
    success = Signal(str)
    failure = Signal(str)

    def __init__(self, filepath, data):
        super().__init__()
        self._filepath = filepath
        self._data     = data

    def run(self):
        try:
            export_to_excel(self._filepath, self._data)
            self.success.emit(self._filepath)
        except Exception as e:
            self.failure.emit(str(e))


class _THDWorker(QThread):
    progress = Signal(str)
    success  = Signal(dict)
    failure  = Signal(str)

    def __init__(self, window_name, num_harmonics):
        super().__init__()
        self._window_name   = window_name
        self._num_harmonics = num_harmonics

    def run(self):
        calc, err = run_thd_calculation(
            self._window_name, self._num_harmonics, self.progress.emit
        )
        if err:
            self.failure.emit(err)
        else:
            self.success.emit(calc)


# ── Main application window ───────────────────────────────────────────────────

class THDApp(QMainWindow):

    def __init__(self):
        super().__init__()
        self.setWindowTitle("THD Analyzer")
        self.setMinimumWidth(500)
        self._last_calc     = None
        self._loader        = None
        self._worker        = None
        self._export_worker = None
        self._build_ui()
        self._start_window_load()

    # ── UI construction ──────────────────────────────────────────────────────

    def _build_ui(self):
        central = QWidget()
        self.setCentralWidget(central)
        root = QVBoxLayout(central)
        root.setContentsMargins(20, 20, 20, 16)
        root.setSpacing(14)

        # Header: logo on the right, title on the left
        header = QHBoxLayout()

        title_col = QVBoxLayout()
        title_col.setSpacing(2)
        lbl_title = QLabel("THD Analyzer")
        lbl_title.setObjectName("AppTitle")
        lbl_sub   = QLabel("Total Harmonic Distortion Calculator")
        lbl_sub.setObjectName("AppSubtitle")
        title_col.addWidget(lbl_title)
        title_col.addWidget(lbl_sub)

        logo_label = QLabel()
        logo_label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        pixmap = QPixmap(_logo_path())
        if not pixmap.isNull():
            logo_label.setPixmap(
                pixmap.scaledToHeight(38, Qt.SmoothTransformation)
            )

        header.addLayout(title_col)
        header.addStretch()
        header.addWidget(logo_label)
        root.addLayout(header)

        # Config frame
        cfg_frame = QFrame()
        cfg_frame.setObjectName("ConfigFrame")
        cfg_grid = QGridLayout(cfg_frame)
        cfg_grid.setContentsMargins(16, 12, 16, 12)
        cfg_grid.setHorizontalSpacing(12)
        cfg_grid.setVerticalSpacing(10)

        lbl_win = QLabel("FFT spectrum window to analyse:")
        lbl_win.setObjectName("ConfigLabel")
        cfg_grid.addWidget(lbl_win, 0, 0, Qt.AlignLeft | Qt.AlignVCenter)

        win_row = QHBoxLayout()
        win_row.setSpacing(6)
        self._window_combo = QComboBox()
        self._window_combo.addItem("(connecting…)")
        self._window_combo.setEnabled(False)

        self._refresh_btn = QPushButton("↺")
        self._refresh_btn.setObjectName("RefreshButton")
        self._refresh_btn.setToolTip("Refresh window list from NVGate")
        self._refresh_btn.clicked.connect(self._start_window_load)

        win_row.addWidget(self._window_combo)
        win_row.addWidget(self._refresh_btn)
        win_row.addStretch()
        cfg_grid.addLayout(win_row, 0, 1)

        lbl_harm = QLabel("Number of harmonics:")
        lbl_harm.setObjectName("ConfigLabel")
        cfg_grid.addWidget(lbl_harm, 1, 0, Qt.AlignLeft | Qt.AlignVCenter)

        self._harmonics_sel = HarmonicsSelector(minimum=2, maximum=20,
                                                initial=DEFAULT_HARMONICS)
        cfg_grid.addWidget(self._harmonics_sel, 1, 1, Qt.AlignLeft)

        root.addWidget(cfg_frame)

        # Result cards — 2 × 2 grid
        cards_grid = QGridLayout()
        cards_grid.setSpacing(12)
        self._card_thd_pct  = ResultCard("THD", "%")
        self._card_thd_db   = ResultCard("THD", "dB")
        self._card_h1_freq  = ResultCard("H1 — Fundamental frequency", "Hz")
        self._card_h1_amp   = ResultCard("H1 — Fundamental amplitude", "V")
        cards_grid.addWidget(self._card_thd_pct,  0, 0)
        cards_grid.addWidget(self._card_thd_db,   0, 1)
        cards_grid.addWidget(self._card_h1_freq,  1, 0)
        cards_grid.addWidget(self._card_h1_amp,   1, 1)
        root.addLayout(cards_grid)

        # Status panel — prominent colored bar above the buttons
        self._status_frame = QFrame()
        self._status_frame.setObjectName("StatusPanel")
        status_row = QHBoxLayout(self._status_frame)
        status_row.setContentsMargins(14, 8, 14, 8)
        status_row.setSpacing(10)

        self._status_dot = QLabel("●")
        self._status_dot.setFixedWidth(14)

        self._status_label = QLabel("Connecting to NVGate…")
        self._status_label.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Preferred)
        self._status_label.setWordWrap(True)

        status_row.addWidget(self._status_dot)
        status_row.addWidget(self._status_label)
        root.addWidget(self._status_frame)

        # Buttons
        btn_row = QHBoxLayout()
        btn_row.setSpacing(12)
        self._run_btn = QPushButton("▶   Run THD")
        self._run_btn.setObjectName("RunButton")
        self._run_btn.setEnabled(False)
        self._run_btn.clicked.connect(self._on_run)

        self._export_btn = QPushButton("Export to Excel")
        self._export_btn.setObjectName("ExportButton")
        self._export_btn.setEnabled(False)
        self._export_btn.clicked.connect(self._on_export)

        btn_row.addWidget(self._run_btn)
        btn_row.addWidget(self._export_btn)
        root.addLayout(btn_row)

        self._set_status("Connecting to NVGate…", "info")

    # ── Status helper ────────────────────────────────────────────────────────

    def _set_status(self, msg, state="info"):
        bg, text, border, dot = _STATUS_STYLES.get(state, _STATUS_STYLES["info"])
        self._status_frame.setStyleSheet(
            f"QFrame#StatusPanel {{"
            f"  background-color: {bg};"
            f"  border: 1px solid {border};"
            f"  border-radius: 6px;"
            f"}}"
        )
        self._status_dot.setStyleSheet(
            f"color: {dot}; background: transparent; font-size: 10px;"
        )
        self._status_label.setStyleSheet(
            f"color: {text}; background: transparent; font-size: 13px;"
        )
        self._status_label.setText(msg)

    # ── Window list loading ──────────────────────────────────────────────────

    def _start_window_load(self):
        self._run_btn.setEnabled(False)
        self._refresh_btn.setEnabled(False)
        self._window_combo.setEnabled(False)
        self._set_status("Connecting to NVGate…", "working")

        self._loader = _WindowLoader()
        self._loader.done.connect(self._on_windows_loaded)
        self._loader.start()

    def _on_windows_loaded(self, windows):
        self._refresh_btn.setEnabled(True)
        if not windows:
            self._window_combo.clear()
            self._window_combo.addItem("(no windows found)")
            self._set_status(
                "Could not reach NVGate — start it in connected mode, then click ↺",
                "error"
            )
            return

        current = self._window_combo.currentText()
        self._window_combo.clear()
        self._window_combo.addItems(windows)
        self._window_combo.setEnabled(True)

        if current in windows:
            self._window_combo.setCurrentText(current)
        elif DEFAULT_WINDOW in windows:
            self._window_combo.setCurrentText(DEFAULT_WINDOW)

        self._run_btn.setEnabled(True)
        self._set_status(f"Ready — {len(windows)} window(s) found", "success")

    # ── Run ──────────────────────────────────────────────────────────────────

    def _on_run(self):
        window        = self._window_combo.currentText()
        num_harmonics = self._harmonics_sel.value()

        self._run_btn.setText("Running…")
        self._run_btn.setEnabled(False)
        self._refresh_btn.setEnabled(False)
        self._export_btn.setEnabled(False)
        self._last_calc = None
        for card in (self._card_thd_pct, self._card_thd_db,
                     self._card_h1_freq, self._card_h1_amp):
            card.set_empty()

        self._set_status(f"Calculating THD on '{window}'…", "working")

        self._worker = _THDWorker(window, num_harmonics)
        self._worker.progress.connect(lambda msg: self._set_status(msg, "working"))
        self._worker.success.connect(self._on_run_success)
        self._worker.failure.connect(self._on_run_failure)
        self._worker.start()

    def _on_run_success(self, calc):
        self._last_calc = calc
        thd    = calc['thd']
        thd_db = calc['thd_db']
        f0     = calc['f0']
        h1     = calc['h1']
        raw_unit = (calc['meta'].get('Unit') or "").strip()
        amp_unit = raw_unit if (raw_unit and raw_unit.upper() != "EU") else "V"

        self._card_thd_pct.set_ready(f"{thd:.3f}", "%")
        self._card_thd_db.set_ready(f"{thd_db:.1f}", "dB")
        self._card_h1_freq.set_ready(f"{f0:.2f}", "Hz")
        self._card_h1_amp.set_ready(_fmt_sci(h1), amp_unit)

        self._run_btn.setText("▶   Run THD")
        self._run_btn.setEnabled(True)
        self._refresh_btn.setEnabled(True)
        self._export_btn.setEnabled(True)
        self._set_status(
            f"Done — THD = {thd:.3f} %  ({thd_db:.1f} dB)", "success"
        )

    def _on_run_failure(self, error):
        self._run_btn.setText("▶   Run THD")
        self._run_btn.setEnabled(True)
        self._refresh_btn.setEnabled(True)
        self._set_status(f"Error: {error}", "error")

    # ── Export ───────────────────────────────────────────────────────────────

    def _on_export(self):
        if self._last_calc is None:
            return

        window_slug  = self._last_calc['window_name'].replace(" ", "_")
        filepath, _ = QFileDialog.getSaveFileName(
            self, "Save Excel report",
            f"THD_{window_slug}.xlsx",
            "Excel workbook (*.xlsx);;All files (*.*)",
        )
        if not filepath:
            return

        self._export_btn.setEnabled(False)
        self._set_status("Writing Excel file…", "working")

        self._export_worker = _ExportWorker(filepath, self._last_calc)
        self._export_worker.success.connect(self._on_export_done)
        self._export_worker.failure.connect(self._on_export_failed)
        self._export_worker.start()

    def _on_export_done(self, filepath):
        self._export_btn.setEnabled(True)
        self._set_status(f"Saved: {filepath}", "success")

    def _on_export_failed(self, error):
        self._export_btn.setEnabled(True)
        self._set_status(f"Export failed: {error}", "error")


# ── Entry point ───────────────────────────────────────────────────────────────

if __name__ == "__main__":
    app = QApplication(sys.argv)
    app.setStyleSheet(_STYLESHEET)
    window = THDApp()
    window.show()
    sys.exit(app.exec())
